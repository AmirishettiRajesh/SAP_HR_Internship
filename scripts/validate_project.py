"""Independently validate generated SAP HR internship deliverables."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = {
    "reports": [
        "Week_1_Workforce_Data_Analysis.docx",
        "Week_2_Payroll_Configuration.docx",
        "Week_3_Performance_Management.docx",
        "Week_4_Succession_Planning.docx",
        "Week_5_HR_Data_Security.docx",
    ],
    "submission_descriptions": [f"Week_{i}_Description.txt" for i in range(1, 6)],
}


def main() -> int:
    errors: list[str] = []
    for folder, names in EXPECTED.items():
        for name in names:
            path = ROOT / folder / name
            if not path.is_file():
                errors.append(f"Missing: {folder}/{name}")
                continue
            if folder == "submission_descriptions":
                count = len(re.findall(r"\b[\w’'-]+\b", path.read_text(encoding="utf-8")))
                print(f"{name}: {count} words")
                if count < 250:
                    errors.append(f"{name}: {count} words; minimum is 250")
            else:
                try:
                    doc = Document(path)
                except Exception as exc:
                    errors.append(f"Unreadable DOCX {name}: {exc}")
                    continue
                body = "\n".join(paragraph.text for paragraph in doc.paragraphs)
                tables = len(doc.tables)
                full_text = "\n".join(
                    [body, *[cell.text for table in doc.tables for row in table.rows for cell in row.cells]]
                )
                words = len(full_text.split())
                images = len(doc.inline_shapes)
                print(f"{name}: {words} document words (paragraphs + tables), {tables} tables, {images} images, {path.stat().st_size:,} bytes")
                for required in ("Executive Summary", "Contents", "References"):
                    if required not in body:
                        errors.append(f"{name}: missing {required}")
                if words < 1800 or tables < 2 or images < 1:
                    errors.append(f"{name}: below structural/content minimum")
                if re.search(r"\b(TODO|TBD|PLACEHOLDER|lorem ipsum|insert here)\b", body, re.I):
                    errors.append(f"{name}: unfinished placeholder detected")

    required_files = [
        "datasets/synthetic_employee_workforce.csv",
        "datasets/synthetic_payroll_scenarios.csv",
        "datasets/synthetic_performance_scores.csv",
        "datasets/synthetic_succession_matrix.csv",
        "charts/workforce_headcount_by_department.png",
        "charts/workforce_headcount_trend.png",
        "charts/workforce_recruitment_source.png",
        "charts/performance_score_distribution.png",
        "charts/succession_nine_box.png",
        "charts/hr_security_risk_matrix.png",
        "diagrams/payroll_process_flow.png",
        "diagrams/succession_governance_cycle.png",
    ]
    for relative in required_files:
        path = ROOT / relative
        if not path.is_file() or path.stat().st_size == 0:
            errors.append(f"Missing or empty: {relative}")

    if errors:
        print("\nVALIDATION FAILED")
        for error in errors:
            print(f"- {error}")
        return 1
    print("\nVALIDATION PASSED: all five DOCX files, descriptions, datasets, and visuals are present.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
